from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from config import CERTIFICATIONS, LOCKED_EDUCATION, LOCKED_EXPERIENCE, LOCKED_IDENTITY, OUTPUT_DIR
from llm_tailor import TailoredResume


def font(run, size, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rf.set(qn("w:eastAsia"), "Times New Roman")


def heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2.2)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(title.upper()), 10.0, True)
    ppr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "5")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "808080")
    pb.append(b)
    ppr.append(pb)


def bullet(doc, text, size=9.0):
    clean = " ".join(text.strip().split())
    if not re.search(r"[.!?]$", clean):
        raise ValueError(f"Refusing to write an unfinished generated bullet: {clean}")
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.09)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.93
    font(p.add_run(clean), size)


def locked_bullet(doc, text, size=9.0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.09)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.93
    font(p.add_run(text.strip()), size)


def tech_header(doc, title, technologies, space_before=0.4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(title), 9.6, True)
    font(p.add_run(" | " + ", ".join(technologies)), 8.5)


def safe(s):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s.strip()).strip("._") or "tailored_resume"


def build_resume(result: TailoredResume) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.28)
    sec.bottom_margin = Inches(0.28)
    sec.left_margin = Inches(0.42)
    sec.right_margin = Inches(0.42)

    # 1. HEADER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0.4)
    font(p.add_run(LOCKED_IDENTITY["display_name"]), 16.0, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    contact = (
        f"{LOCKED_IDENTITY['phone']} | {LOCKED_IDENTITY['email']} | "
        f"LinkedIn: {LOCKED_IDENTITY['linkedin']} | GitHub: {LOCKED_IDENTITY['github']} | "
        f"{LOCKED_IDENTITY['location']}"
    )
    font(p.add_run(contact), 8.35)

    # 2. PROFESSIONAL SUMMARY
    heading(doc, "Professional Summary")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = 0.92
    font(p.add_run(" ".join(result.tailored_summary.strip().split())), 9.05)

    # 3. WORK EXPERIENCE
    heading(doc, "Work Experience")
    x = LOCKED_EXPERIENCE
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(x["company"]), 9.5, True)
    font(p.add_run(f" | {x['role']} | {x['dates']}"), 8.8)
    for b in x["bullets"]:
        locked_bullet(doc, b, size=9.0)

    # 4. INDIVIDUAL PROJECTS
    heading(doc, "Individual Projects")
    for idx, proj in enumerate(result.tailored_projects):
        tech_header(doc, proj.title, proj.technologies, space_before=(0.8 if idx == 0 else 5.0))
        for b in proj.bullets:
            bullet(doc, b)

    # 5. EDUCATION
    heading(doc, "Education")
    e = LOCKED_EDUCATION
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(e["institution"]), 8.95, True)
    font(p.add_run(f" | {e['dates']}"), 8.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(e["degree"]), 8.75, True)
    font(p.add_run(f" | CGPA: {e['cgpa']}"), 8.5)

    # 6. CERTIFICATIONS AND ACHIEVEMENTS
    # Certifications are locked factual entries, not generated prose bullets.
    # They are allowed to omit terminal punctuation and must not be rejected
    # by the generated-bullet completion validator.
    heading(doc, "Certifications and Achievements")
    for cert in CERTIFICATIONS:
        locked_bullet(doc, cert, size=8.75)

    # 7. SKILLS
    heading(doc, "Skills")
    for group in result.tailored_skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 0.9
        font(p.add_run(f"{group.category}: "), 8.5, True)
        font(p.add_run(", ".join(group.items)), 8.5)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True

    path = OUTPUT_DIR / f"{safe(result.company_name)}_{safe(result.job_role)}_tailored_resume.docx"
    doc.save(path)
    return path
